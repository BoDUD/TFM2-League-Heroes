"""Hero documentation generated from the repo's current data (no network).

    python tools/docs/gen_hero_docs.py                       # docs/heroes/<id>.md, the hero table in README.md, docs/heroes.xlsx
    python tools/docs/gen_hero_docs.py --released <file>     # also the 「工坊已发布」 column (one hero id a line, e.g. ahri)
    python tools/docs/gen_hero_docs.py --check               # exit 1 when a page or the README table is out of date

Reads (nothing else, nothing written outside docs/ and README.md):
  league/champion/*.data_champion, tft/champion/*.data_champion   the heroes, their cooldowns, ranges and stats
  league/text/champion.i18n, tft/text/champion.i18n               names, skill names and the in-game skill texts (zh-hans)
  addons/league_main/src/hd.rs                                    PPU (HEROES) and drawn size (SIZES)
  addons/*/mod.mod_info, mod.override_info, README.md, text/      the add-ons: which hero, name, what changes
  docs/preview/<id>_*.gif / .png                                  showcase GIFs and frame sheets (links only)
  league/champions/<id>#sheet.png + #anim.fanim                   the avatar (idle frame 0) in docs/heroes/img/, needs pillow

Each page has two parts: the top is generated (overwritten every run); everything from the line
`## 开发记录` down is hand-written and kept as it is (the old README's section of the hero was moved there once with
--import-readme). The README keeps its hand-written text; only the part between the HEROES markers is generated.

Files missing from the working tree (a sparse checkout) are read from git's HEAD instead, so the output is the same.
The xlsx needs openpyxl (pip install openpyxl; not a dependency of the mod) and is skipped without it; --check does not
look at the xlsx or the avatars (binary).
"""
import argparse
import datetime
import glob
import io
import json
import os
import re
import subprocess
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
PAGES = "docs/heroes"
IMG = "docs/heroes/img"
XLSX = "docs/heroes.xlsx"
REPO_URL = "https://github.com/BoDUD/TFM2-League-Heroes/blob/main/"
MARK_DEV = "## 开发记录"
GEN_NOTE = ("<!-- 本页上半部分由 tools/docs/gen_hero_docs.py 按仓库里现在的数据生成，手改会被覆盖；"
            "「开发记录」以下是手写的，生成时原样保留。 -->")
BEGIN, END = "<!-- HEROES:BEGIN -->", "<!-- HEROES:END -->"

LANES = ["上单", "打野", "中单", "ADC", "辅助"]
LANE_WORDS = {"上单": "上单", "上路": "上单", "打野": "打野", "中单": "中单", "中路": "中单", "ADC": "ADC",
              "下路": "ADC", "辅助": "辅助"}
# heroes whose 开发记录 has no 「定位」 row with a lane (the first ones, written before the table had it)
LANE_FALLBACK = {"league_garen": "上单", "league_ashe": "ADC", "league_lux": "中单"}
# old README section titles that are not the hero's name in league/text
TITLE_ALIAS = {"蛮王": "league_tryndamere", "爆爆": "tft_powder"}

CATEGORY = {"Melee": "近战", "Magician": "法师", "Range": "射手", "Util": "功能型", "Assassin": "刺客"}
TARGET = {"EnemyWithoutTower": "敌人（不含防御塔）", "EnemyChampion": "敌方英雄", "Enemy": "敌人",
          "EnemyChampionRecentlyAttacked": "刚交过手的敌方英雄", "AllyChampion": "友方英雄",
          "AllyNotSelf": "友方（不含自己）", "EnemyChampionInCC": "被控制的敌方英雄", "Self": "自己"}
CAST = {"Targeting": "指定目标", "Direction": "方向", "Position": "地点", "None": "无目标"}
STATS = [("attack", "攻击力"), ("magic_power", "法强"), ("hp", "生命"), ("defence", "护甲"),
         ("magic_resistance", "魔抗"), ("move_speed", "移速"), ("hp_regen", "生命回复"), ("crit_chance", "暴击")]
# stat icons in the skill texts: the word shown when the text itself does not say it right after the number
ICON_WORD = {"ad_0": "攻击力", "ap_0": "法强", "hp_0": "最大生命", "speed_0": "移速", "attack_speed_0": "攻速",
             "armor_0": "护甲", "range_0": "射程", "mr_0": "魔抗"}
ICON_SAID = {"ad_0": ("攻击",), "ap_0": ("法术强度", "法强"), "hp_0": ("生命",), "speed_0": ("移动速度", "移速", "速度"),
             "attack_speed_0": ("攻速", "攻击速度"), "armor_0": ("护甲",), "range_0": ("射程", "距离", "范围"),
             "mr_0": ("魔抗", "魔法抗性")}
SLOTS = [("attack", "普攻 / 被动"), ("skill", "技能1"), ("skill2", "技能2"), ("ult", "大招")]
SLOT_NAME = {"skill": "skill1", "skill2": "skill2", "ult": "ult"}
TICKS = 60  # ticks a second
# add-ons that are no hero's (the DLL, the bundle, shared code, diagnostics)
NOT_HERO_ADDONS = {"league_main", "league_addons", "game_places", "vendor", "league_recall_log", "league_unlock_guard"}


# ---------------------------------------------------------------- reading the repo (working tree, else git HEAD)
_git_files = None


def git_files():
    global _git_files
    if _git_files is None:
        try:
            out = subprocess.run(["git", "-C", ROOT, "ls-files", "-z"], capture_output=True, check=True).stdout
            _git_files = set(out.decode("utf-8").split("\0")) - {""}
        except (OSError, subprocess.CalledProcessError):
            _git_files = set()
    return _git_files


def exists(rel):
    return os.path.exists(os.path.join(ROOT, rel)) or rel in git_files()


def dir_exists(rel):
    rel = rel.rstrip("/") + "/"
    return os.path.isdir(os.path.join(ROOT, rel)) or any(f.startswith(rel) for f in git_files())


def read_bytes(rel):
    p = os.path.join(ROOT, rel)
    if os.path.exists(p):
        with open(p, "rb") as f:
            return f.read()
    if rel in git_files():
        return subprocess.run(["git", "-C", ROOT, "show", "HEAD:" + rel], capture_output=True, check=True).stdout
    return None


def read_text(rel):
    b = read_bytes(rel)
    return None if b is None else b.decode("utf-8-sig").replace("\r\n", "\n")


def read_json(rel):
    t = read_text(rel)
    return None if t is None else json.loads(t)


def list_files(pattern):
    """Repo-relative paths matching a glob, from the working tree and git's index."""
    found = {os.path.relpath(p, ROOT).replace("\\", "/") for p in glob.glob(os.path.join(ROOT, pattern))}
    rx = re.compile("^" + re.escape(pattern).replace(r"\*", "[^/]*") + "$")
    found |= {f for f in git_files() if rx.match(f)}
    return sorted(found)


# ---------------------------------------------------------------- text
def plain(s):
    """In-game rich text to plain text: colours dropped, a stat icon said in words where the text does not say it."""
    s = s.replace("\r\n", "\n")
    s = re.sub(r"<i#[^>]*:(\w+)>", lambda m: "\x01" + m.group(1) + "\x02", s)
    s = re.sub(r"<#[0-9a-fA-F]{6,8}>", "", s).replace("<>", "")
    s = re.sub(r"<[^<>\n]{0,80}>", "", s)
    out, i = [], 0
    for m in re.finditer(r"\x01(\w+)\x02", s):
        out.append(s[i:m.start()])
        i = m.end()
        key = m.group(1)
        num = re.match(r"\s*[\d.,/]+%?", s[i:])
        if num:
            after = s[i + num.end():].lstrip()
            after = re.sub(r"\x01\w+\x02", "", after)
            after = re.split(r"[，。；、,.;：:（）()+\-\d\s]", after, maxsplit=1)[0][:6]  # the words right after it
            if not any(w in after for w in ICON_SAID.get(key, ())):
                out.append(num.group(0) + " " + ICON_WORD.get(key, key))
                i += num.end()
    out.append(s[i:])
    s = "".join(out)
    s = re.sub(r"\x01\w+\x02", "", s)
    return s.strip()


def md_cell(s):
    return s.replace("|", "\\|").replace("\n", "<br>")


def md_para(s):
    return s.replace("\n", "  \n")


def strip_md(s):
    s = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", s)
    s = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", s)
    return s.replace("**", "").replace("`", "").strip()


def short(s, n):
    s = re.sub(r"\s+", " ", s).strip()
    return s if len(s) <= n else s[:n - 1] + "…"


def rel_link(target, from_dir):
    """A repo-relative path (maybe with ?query / #anchor) as seen from from_dir (repo-relative)."""
    m = re.match(r"([^?#]*)(.*)", target)
    path, rest = m.group(1), m.group(2)
    if not path:
        return target
    return os.path.relpath(path, from_dir or ".").replace("\\", "/") + rest


def rewrite_links(text, src_dir, dst_dir, anchors=None):
    """Relative links in markdown written in src_dir, rewritten to resolve from dst_dir. `anchors` maps a bare
    #anchor of the old README to the page it now lives on (repo-relative)."""
    def fix(m):
        bang, label, target = m.group(1), m.group(2), m.group(3)
        if re.match(r"[a-z]+:", target) or target.startswith("/"):
            return m.group(0)
        if target.startswith("#"):
            page = (anchors or {}).get(target[1:])
            if page is None:
                return m.group(0)
            return "%s[%s](%s)" % (bang, label, rel_link(page + target, dst_dir))
        full = os.path.normpath(os.path.join(src_dir or ".", re.match(r"[^?#]*", target).group(0))).replace("\\", "/")
        rest = target[len(re.match(r"[^?#]*", target).group(0)):]
        return "%s[%s](%s)" % (bang, label, rel_link(full + rest, dst_dir))
    return re.sub(r"(!?)\[((?:[^\[\]]|\[[^\]]*\])*)\]\(([^()\s]*(?:\([^()]*\)[^()\s]*)*)\)", fix, text)


# ---------------------------------------------------------------- the data
def load_texts():
    zh, en = {}, {}
    for f in ("league/text/champion.i18n", "tft/text/champion.i18n"):
        d = read_json(f)
        for hid, v in d["zh-hans"]["description"].items():
            zh[hid] = dict(v)
        for hid, v in d["zh-hans"].get("skill_name", {}).items():
            zh.setdefault(hid, {})["_names"] = v
        for hid, v in d.get("en", {}).get("description", {}).items():
            en[hid] = v.get("name", "")
    return zh, en


def load_ppu():
    src = read_text("addons/league_main/src/hd.rs") or ""
    ppu, size = {}, {}
    m = re.search(r"pub const HEROES[^=]*=\s*&\[(.*?)\n\];", src, re.S)
    if m:
        for hid, p in re.findall(r'\("((?:league|tft)_\w+)",\s*([\d.]+),\s*include_str', m.group(1)):
            ppu[hid] = float(p)
    m = re.search(r"pub const SIZES[^=]*=\s*&\[(.*?)\];", src, re.S)
    if m:
        for hid, s in re.findall(r'\("((?:league|tft)_\w+)",\s*([\d.]+)\)', m.group(1)):
            size[hid] = float(s)
    return ppu, size


def fmt_num(x):
    return ("%.2f" % x).rstrip("0").rstrip(".")


def readme_tables(md, words):
    """The add-on README's tables whose header has one of `words`: [(header cells, [row cells])]."""
    lines, out = md.split("\n"), []
    for i, l in enumerate(lines):
        if l.startswith("|") and any(w in l for w in words) and i + 1 < len(lines) and re.match(r"\|\s*-", lines[i + 1]):
            head = [c.strip() for c in l.strip().strip("|").split("|")]
            rows = []
            for r in lines[i + 2:]:
                if not r.startswith("|"):
                    break
                rows.append([c.strip() for c in re.split(r"(?<!\\)\|", r.strip().strip("|"))])
            out.append((head, rows))
    return out


def addon_name(heading, hero_zh):
    h = heading.lstrip("#").strip()
    h = re.sub(r"[（(][^）)]*v\d[^）)]*[）)]", "", h)
    h = re.sub(r"[（(]`[^`]*`[）)]", "", h).strip()
    if not re.search(r"[\u4e00-\u9fff]", h):
        h = "%s附加包" % hero_zh
    return h


def first_para(md):
    """The add-on README's first paragraph that says what it does (markdown), with the list it introduces."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", md)]
    for i, p in enumerate(paras[1:], 1):
        if not p or p.startswith(("#", "|", "```", "!")) or p.startswith("用户"):
            continue
        if p.endswith(("：", ":")) and i + 1 < len(paras) and paras[i + 1].startswith("- "):
            p += "\n\n" + paras[i + 1]
        return p
    return ""


def load_addons(heroes, zh):
    """{hero id: [add-on dict]} and the add-ons that are no single hero's."""
    names = sorted(((zh[h]["name"], h) for h in heroes), key=lambda x: -len(x[0]))

    def hero_of_cell(cell):
        cell = strip_md(cell)
        hit = []
        for part in re.split(r"[、/，,]\s*", cell):
            for n, h in names:
                if part.startswith(n):
                    hit.append(h)
                    break
        return hit

    per_hero, general = {}, []
    for mi in list_files("addons/*/mod.mod_info"):
        d = os.path.dirname(mi)
        aid = os.path.basename(d)
        if aid in NOT_HERO_ADDONS:
            continue
        info = read_json(mi)
        over = read_json(d + "/mod.override_info") or {}
        targets = [k.split("/")[-1] for k in over if "/champion/" in k]
        targets = [t for t in targets if t in heroes]
        if not targets and aid in heroes:
            targets = [aid]
        md = read_text(d + "/README.md") or ""
        generic = not targets
        # a hero's add-on: its main-pack-vs-add-on table; one for many heroes (league_ult_press, league_ult_aim):
        # those and the 「什么时候按」 tables, only the rows naming the hero
        tables = readme_tables(md, ("主包", "什么时候") if generic else ("主包",))
        if not generic:
            tables = tables[:1]
        by_row = generic or len(targets) > 1 or any(t[0] and t[0][0] == "英雄" for t in tables)
        tables = [(head, [(r, hero_of_cell(r[0])) for r in rows]) for head, rows in tables]
        if generic:
            targets = sorted({h for _, rows in tables for _, hs in rows for h in hs})
        heading = md.split("\n", 1)[0] if md.startswith("#") else ""
        texts = (read_json(d + "/text/champion.i18n") or {}).get("zh-hans", {}).get("description", {})
        for h in targets:
            mine = [(head, [r for r, hs in rows if h in hs or not by_row]) for head, rows in tables]
            mine = [(head, rows) for head, rows in mine if rows]
            a = {"id": aid, "dir": d, "version": info.get("version", ""), "generic": generic,
                 "name": addon_name(heading, zh[h]["name"]) if heading else "%s附加包" % zh[h]["name"],
                 "tables": mine, "para": first_para(md), "desc": info.get("description", ""),
                 "texts": texts.get(h, {}), "has_readme": bool(md)}
            if generic:
                a["name"] = addon_name(heading, zh[h]["name"]) if heading else aid
            per_hero.setdefault(h, []).append(a)
        if generic:
            general.append(aid)
    for v in per_hero.values():
        v.sort(key=lambda a: (a["generic"], a["id"]))
    return per_hero, general


def addon_summary(a, n=90):
    rows = [r for _, rs in a["tables"] for r in rs]
    if rows:
        parts = []
        for r in rows:
            label = strip_md(r[0]) if r[0] else ""
            val = strip_md(r[-1]) if len(r) > 1 else ""
            parts.append("%s：%s" % (label, short(val, n)) if label else short(val, n))
        return "\n".join(parts)
    if a["para"]:
        return short(strip_md(a["para"]), 2 * n)
    return short(a["desc"], 2 * n)


# ---------------------------------------------------------------- the old README (--import-readme, once)
def split_old_readme(text):
    """(intro lines, gallery lines, [(title, lines)] of every '## ' section)."""
    lines = text.replace("\r\n", "\n").split("\n")
    first = next(i for i, l in enumerate(lines) if l.startswith("## "))
    top = lines[1:first]
    # the gallery: from the first hero showcase / combo GIF (the idle-breath GIF before it belongs to its paragraph)
    g = next((i for i, l in enumerate(top) if l.startswith("![") and re.search(r"_(showcase|combos)\.gif", l)),
             len(top))
    sections, cur = [], None
    for l in lines[first:]:
        if l.startswith("## "):
            cur = (l[3:].strip(), [l])
            sections.append(cur)
        else:
            cur[1].append(l)
    return top[:g], top[g:], sections


def gh_anchor(title):
    t = title.strip().lower()
    t = re.sub(r"[^\w\- \u4e00-\u9fff]", "", t)
    return t.replace(" ", "-")


def hero_of_title(title, name2id):
    base = re.split(r"[（(]", title[len("英雄："):])[0].strip()
    return name2id.get(base) or TITLE_ALIAS.get(base)


# ---------------------------------------------------------------- avatars
def make_avatar(hid):
    try:
        from PIL import Image
    except ImportError:
        return None
    folder = "tft/champions" if hid.startswith("tft_") else "league/champions"
    sheet = read_bytes("%s/%s#sheet.png" % (folder, hid))
    anim = read_text("%s/%s#anim.fanim" % (folder, hid))
    if sheet is None or anim is None:
        return None
    fr = json.loads(anim)["anims"]["idle"]["frames"][0]["data"]
    im = Image.open(io.BytesIO(sheet)).convert("RGBA")
    x, y, w, h = (int(fr[k]) for k in ("x", "y", "w", "h"))
    im = im.crop((x, y, x + w, y + h))
    box = im.getbbox()
    if box:
        im = im.crop(box)
    side = max(im.size)
    sq = Image.new("RGBA", (side, side), (0, 0, 0, 0))
    sq.paste(im, ((side - im.size[0]) // 2, side - im.size[1]))
    k = max(1, 64 // side)
    if k > 1:
        sq = sq.resize((side * k, side * k), Image.NEAREST)
    out = io.BytesIO()
    sq.save(out, "PNG", optimize=True)
    return out.getvalue()


# ---------------------------------------------------------------- one hero
class Hero:
    pass


def build_heroes(args):
    zh, en = load_texts()
    ppu, size = load_ppu()
    ids = [os.path.basename(p)[:-len(".data_champion")] for p in list_files("league/champion/*.data_champion")]
    ids += [os.path.basename(p)[:-len(".data_champion")] for p in list_files("tft/champion/*.data_champion")]
    heroes = {}
    for hid in ids:
        folder = "tft" if hid.startswith("tft_") else "league"
        h = Hero()
        h.id, h.short = hid, hid.split("_", 1)[1]
        h.tft = hid.startswith("tft_")
        h.data = read_json("%s/champion/%s.data_champion" % (folder, hid))
        t = zh.get(hid, {})
        h.zh, h.en = t.get("name", hid), en.get(hid, "")
        h.text = {s: plain(t.get(s, "")) for s, _ in SLOTS}
        h.names = t.get("_names", {})
        h.ppu, h.size = ppu.get(hid), size.get(hid)
        h.page = "%s/%s.md" % (PAGES, hid)
        gifs = list_files("docs/preview/%s_*.gif" % hid)
        h.showcase = "docs/preview/%s_showcase.gif" % hid if "docs/preview/%s_showcase.gif" % hid in gifs else None
        h.gifs = [g for g in gifs if g != h.showcase]
        h.pngs = [p for p in ("docs/preview/%s_frames.png" % hid, "docs/preview/%s_effects.png" % hid) if exists(p)]
        heroes[hid] = h
    addons, general = load_addons(set(heroes), zh)
    for hid, h in heroes.items():
        h.addons = addons.get(hid, [])
        h.passive = None
        m = re.search(r"被动\s*「?([^：:「」，。\s]{1,12})」?\s*[：:]", h.text["attack"])
        if m:
            h.passive = m.group(1)
    return heroes, general, zh


def lane_of(dev_text, hid):
    for l in dev_text.split("\n"):
        if re.match(r"\|\s*定位\s*\|", l):
            m = re.search("|".join(LANE_WORDS), l)
            if m:
                return LANE_WORDS[m.group(0)]
            break
    if hid == "tft_powder":
        return "中单"
    return LANE_FALLBACK.get(hid, "")


def skill_title(h, slot, label):
    if slot == "attack":
        return label + ("（被动「%s」）" % h.passive if h.passive else "")
    n = h.names.get(SLOT_NAME[slot])
    return "%s · %s" % (label, n) if n else label


def skill_text(h, slot):
    s = h.text[slot]
    n = h.names.get(SLOT_NAME.get(slot, ""), "")
    if n and s.startswith(n) and s[len(n):len(n) + 1] in ("：", ":"):
        s = s[len(n) + 1:].lstrip()
    return s


def skill_meta(h, slot):
    d = h.data[slot]
    parts = []
    cd = d.get("cooltime", 0) / TICKS
    if slot == "attack":
        parts.append("攻击间隔 %s 秒" % fmt_num(cd))
        parts.append("攻击距离 %d" % d.get("range", 0))
    else:
        parts.append("冷却 %s 秒" % fmt_num(cd))
        if d.get("cooltime_use_count"):
            parts.append("可存 %d 次" % d["cooltime_use_count"])
        parts.append("施放距离 %d" % d.get("range", 0))
        tgt = TARGET.get(d.get("casting_target"), d.get("casting_target", ""))
        parts.append("施放：%s · %s" % (CAST.get(d.get("casting_type"), d.get("casting_type", "")), tgt))
    return " · ".join(parts)


def ppu_text(h):
    if not h.ppu:
        return "1 倍像素（无高清帧）"
    s = "PPU %s" % fmt_num(h.ppu)
    if h.size:
        s += "（画成 %d%%）" % round(h.size * 100)
    return s


def one_line(h):
    parts = []
    if h.passive:
        parts.append("被动 %s" % h.passive)
    for slot, label in (("skill", "技能1"), ("skill2", "技能2"), ("ult", "大招")):
        n = h.names.get(SLOT_NAME[slot])
        if n:
            parts.append(("大招 " if slot == "ult" else "") + n)
    return " · ".join(parts)


def related_files(h):
    s = h.short
    cand = ["tools/kit/build_%s.py" % s, "tools/kit/setup_%s.py" % s, "tools/kit/%s_combos.py" % s,
            "tools/lol/extract_%s.py" % s, "tools/art/preview_%s.py" % s, "tools/art/rig_%s.py" % s,
            "tools/art/import_%s.py" % s, "tools/art/design_%s.py" % s]
    files = [c for c in cand if exists(c)]
    if dir_exists("assets/source/%s" % s):
        files.append("assets/source/%s/" % s)
    return files


def page_top(h, avatar_ok):
    L = []
    title = "# %s %s" % (h.zh, h.en) if h.en and h.en != h.zh else "# %s" % h.zh
    L += [GEN_NOTE, "", title, ""]
    if avatar_ok:
        L += ['<img src="img/%s.png" alt="%s" width="64">' % (h.id, h.zh), ""]
    d = h.data
    rows = [("id", "`%s`" % h.id),
            ("所属的包", "tft 包（金铲铲之战系列，mod_id `tft`）" if h.tft else "主包（mod_id `league`）"),
            ("分路", h.lane or "（未标注）"),
            ("分类 / 标签", "%s（%s） · %s" % (CATEGORY.get(d["category"], d["category"]), d["category"],
                                         " / ".join(d.get("tags", [])))),
            ("高清帧", ppu_text(h))]
    if h.addons:
        rows.append(("附加包", "、".join("[%s](#附加包)（`%s`）" % (a["name"], a["id"]) for a in h.addons)))
    else:
        rows.append(("附加包", "无"))
    L += ["| 项目 | 内容 |", "|---|---|"] + ["| %s | %s |" % (k, md_cell(v)) for k, v in rows] + [""]
    if h.showcase:
        L += ["![%s演示](%s)" % (h.zh, rel_link(h.showcase, PAGES)), ""]
    more = ["[%s](%s)" % (os.path.basename(g)[len(h.id) + 1:-4], rel_link(g, PAGES)) for g in h.gifs]
    more += ["[%s](%s)" % ("逐帧预览" if p.endswith("_frames.png") else "特效表", rel_link(p, PAGES)) for p in h.pngs]
    if more:
        L += ["其他预览：" + "、".join(more), ""]
    L += ["## 现在的技能", "",
          "下面是游戏里显示的技能说明（`%s/text/champion.i18n` 的简体中文，去掉了颜色和图标），冷却、距离来自 "
          "`%s/champion/%s.data_champion`；1 秒 = %d tick。" % ("tft" if h.tft else "league", "tft" if h.tft else "league",
                                                                 h.id, TICKS), ""]
    for slot, label in SLOTS:
        L += ["### %s" % skill_title(h, slot, label), "", "*%s*" % skill_meta(h, slot), "",
              md_para(skill_text(h, slot)) or "（没有说明）", ""]
    st, gr = d.get("stat", {}), d.get("growth", {})
    parts = []
    for k, n in STATS:
        if st.get(k) or gr.get(k):
            parts.append("%s %s（每级 +%s）" % (n, fmt_num(st.get(k, 0)), fmt_num(gr.get(k, 0))))
    L += ["### 数值", "", "、".join(parts) + "。", ""]
    if h.addons:
        L += ["## 附加包", "",
              ("附加包是可选的原生代码 mod（tft 包的附加包单独安装，不在附加包合集 `league_addons` 里），**不装 tft 包里的英雄"
               "也是完整的**。" if h.tft else
               "附加包是可选的原生代码 mod（都在附加包合集 `league_addons` 里），**不装主包里的英雄也是完整的**。"), ""]
        for a in h.addons:
            link = "%s/README.md" % a["dir"]
            name = "[%s](%s)" % (a["name"], rel_link(link, PAGES)) if a["has_readme"] else a["name"]
            L += ["### %s" % a["name"], "",
                  "%s：`%s`%s%s" % (name, a["id"], " v%s" % a["version"] if a["version"] else "",
                                   "（通用附加包，下面只列和%s有关的部分）" % h.zh if a["generic"] else ""), ""]
            for head, rows in a["tables"]:
                L += ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
                for r in rows:
                    r = (r + [""] * len(head))[:len(head)]
                    L.append("| " + " | ".join(rewrite_links(c, a["dir"], PAGES) for c in r) + " |")
                L.append("")
            if not a["tables"] and a["para"]:
                L += [rewrite_links(a["para"], a["dir"], PAGES), ""]
            elif not a["tables"] and a["desc"]:
                L += ["（这个附加包没有 README，下面是它 `mod.mod_info` 里的英文说明）", "", "> " + a["desc"], ""]
            changed = [(s, lab) for s, lab in SLOTS
                       if a["texts"].get(s) and plain(a["texts"][s]) != h.text[s]]
            if changed:
                L += ["装上附加包后游戏里显示的说明（和主包不同的部分）：", ""]
                for s, lab in changed:
                    L += ["- **%s**：%s" % (skill_title(h, s, lab), plain(a["texts"][s]).replace("\n", " "))]
                L.append("")
    files = related_files(h)
    if files:
        L += ["## 相关脚本和素材", ""] + ["- [`%s`](%s)" % (f, rel_link(f, PAGES)) for f in files] + [""]
    return "\n".join(L)


def split_page(text):
    if text is None:
        return None
    i = text.find("\n" + MARK_DEV + "\n")
    if i < 0:
        return None
    return text[i + 1:]


DEV_EMPTY = MARK_DEV + "\n\n（还没有记录。这个英雄的设计、美术、音效和改动的说明写在这里。）\n"


# ---------------------------------------------------------------- README table, xlsx
def readme_table_md(heroes):
    L = [BEGIN, "<!-- 这一段由 tools/docs/gen_hero_docs.py 生成，请不要手改。 -->", ""]
    total = len(heroes)
    n_league = sum(1 for h in heroes.values() if not h.tft)
    L += ["共 %d 位：主包 %d 位、tft 包 %d 位。点名字看这位英雄现在的技能、附加包和开发记录；表格版见 [docs/heroes.xlsx](docs/heroes.xlsx)。" %
          (total, n_league, total - n_league), ""]
    order = LANES + [""]
    for lane in order:
        hs = sorted((h for h in heroes.values() if h.lane == lane), key=lambda h: (h.tft, h.en.lower(), h.id))
        if not hs:
            continue
        L += ["### %s（%d）" % (lane or "未标注分路", len(hs)), "",
              "| 头像 | 英雄 | 分路 | 技能 | 附加包 | 高清帧 |", "|---|---|---|---|---|---|"]
        for h in hs:
            img = '<img src="docs/heroes/img/%s.png" width="32">' % h.id if h.avatar else ""
            name = "[%s](docs/heroes/%s.md)<br>%s" % (h.zh, h.id, h.en)
            if h.tft:
                name += "（tft 包）"
            ad = "<br>".join("[%s](addons/%s/README.md)" % (a["name"], a["id"]) if a["has_readme"] else a["name"]
                             for a in h.addons if not a["generic"]) or "—"
            pp = "PPU %s" % fmt_num(h.ppu) if h.ppu else "—"
            L.append("| %s | %s | %s | %s | %s | %s |" % (img, name, h.lane or "—", md_cell(one_line(h)), ad, pp))
        L.append("")
    L.append(END)
    return "\n".join(L)


def write_xlsx(heroes, released, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed: %s skipped (pip install openpyxl)" % path)
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "英雄"
    cols = [("中文名", 10), ("English", 14), ("id", 20), ("分路", 7), ("普攻/被动", 50), ("技能1（名+说明）", 60),
            ("技能2", 60), ("大招", 60), ("附加包（名称+作用）", 55), ("PPU", 9), ("工坊已发布", 9), ("文档链接", 30)]
    ws.append([c for c, _ in cols])
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    order = {l: i for i, l in enumerate(LANES)}
    for h in sorted(heroes.values(), key=lambda h: (h.tft, order.get(h.lane, 9), h.en.lower())):
        def sk(slot):
            n = h.names.get(SLOT_NAME.get(slot, ""), "")
            body = skill_text(h, slot)
            return ("%s：%s" % (n, body) if n else body) + "\n（%s）" % skill_meta(h, slot)
        ad = "\n\n".join("%s（%s%s）\n%s" % (a["name"], a["id"], " v" + a["version"] if a["version"] else "",
                                              addon_summary(a)) for a in h.addons)
        if released is None:
            rel = ""
        elif h.tft:
            rel = "—（tft 包）"
        else:
            rel = "是" if h.short in released else "否"
        url = REPO_URL + h.page
        ws.append([h.zh, h.en, h.id, h.lane, sk("attack"), sk("skill"), sk("skill2"), sk("ult"), ad,
                   fmt_num(h.ppu) if h.ppu else "—", rel, url])
        c = ws.cell(row=ws.max_row, column=len(cols))
        c.hyperlink = url
        c.font = Font(color="0563C1", underline="single")
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    when = datetime.datetime(2026, 1, 1)
    wb.properties.creator = "tools/docs/gen_hero_docs.py"
    wb.properties.created = when
    wb.properties.modified = when
    os.makedirs(os.path.dirname(os.path.join(ROOT, path)) or ".", exist_ok=True)
    wb.save(os.path.join(ROOT, path))
    return True


# ---------------------------------------------------------------- import of the old README
def import_readme(path, heroes, zh):
    """Once: every '## 英雄：' section of the old README into its hero's 开发记录; the intro, the gallery and the other
    sections into docs/ pages. Returns {repo-relative file: text}."""
    with open(path, encoding="utf-8") as f:
        text = f.read()
    intro, gallery, sections = split_old_readme(text)
    name2id = {zh[h]["name"]: h for h in heroes}
    keep_in_readme = {"安装测试", "仓库里的 skill", "快速开始", "知识来源", "声明"}
    anchors, out, hero_secs, notes = {}, {}, {}, []
    for title, lines in sections:
        if title.startswith("英雄："):
            hid = hero_of_title(title, name2id)
            if hid is None:
                raise SystemExit("cannot tell the hero of section %r" % title)
            if hid in hero_secs:
                raise SystemExit("two sections for %s" % hid)
            hero_secs[hid] = (title, lines)
            anchors[gh_anchor(title)] = "%s/%s.md" % (PAGES, hid)
        elif title not in keep_in_readme:
            notes.append((title, lines))
            anchors[gh_anchor(title)] = "docs/dev-notes.md"
    for hid, (title, lines) in hero_secs.items():
        body = "\n".join(lines[1:]).strip("\n")
        body = rewrite_links(body, "", PAGES, anchors)
        out["%s/%s.md" % (PAGES, hid)] = (MARK_DEV + "\n\n> 下面是原来 README 里「%s」一节，原样搬过来（只改了相对链接）。"
                                          "\n\n%s\n" % (title, body))
    # the intro paragraphs, each under a heading of its own
    T = ["# 技术说明（原 README 开头）", "",
         "原来 README 开头的各段说明，原样搬到这里（只改了相对链接、加了小标题）。", ""]
    special = {0: "项目简介"}
    for i, para in enumerate(p for p in intro if p.strip()):
        if para.startswith("!["):
            T += [rewrite_links(para, "", "docs", anchors), ""]
            continue
        if para.startswith("英雄："):
            head = "英雄名单和分组"
        elif para.startswith("对敌人放的小技能"):
            head = "小技能的施放目标"
        else:
            m = re.match(r"\**([^（：*(]{1,24})", para)
            head = special.get(i) or (m.group(1).strip() if m else "说明")
        T += ["## " + head, "", rewrite_links(para, "", "docs", anchors), ""]
    out["docs/tech-notes.md"] = "\n".join(T)
    G = ["# 演示动图（原 README 开头）", "",
         "原来 README 开头的全部演示动图，原样搬到这里。每位英雄的动图也在 [英雄页](heroes/) 里。", ""]
    G += [rewrite_links(l, "", "docs", anchors) for l in gallery]
    out["docs/gallery.md"] = "\n".join(G).rstrip("\n") + "\n"
    N = ["# 开发笔记（原 README 里不属于某一位英雄的几节）", ""]
    for title, lines in notes:
        N += [rewrite_links("\n".join(lines).strip("\n"), "", "docs", anchors), ""]
    out["docs/dev-notes.md"] = "\n".join(N)
    return out, hero_secs, anchors


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--check", action="store_true", help="only compare; exit 1 when something is out of date")
    ap.add_argument("--released", help="the hero ids on the Workshop, one a line (fills 「工坊已发布」 in the xlsx)")
    ap.add_argument("--xlsx", default=XLSX, help="where the spreadsheet goes (default %(default)s); '' for none")
    ap.add_argument("--import-readme", help="once: move the hero sections of this old README into the pages")
    args = ap.parse_args()

    heroes, general, zh = build_heroes(args)
    imported = {}
    if args.import_readme:
        imported, _, _ = import_readme(args.import_readme, heroes, zh)

    wanted = {}
    for hid, h in sorted(heroes.items()):
        old = read_text(h.page) if not args.import_readme else None
        dev = imported.get(h.page) or split_page(old) or DEV_EMPTY
        h.lane = lane_of(dev, hid)
        h.dev = dev
        h.avatar = exists("%s/%s.png" % (IMG, hid))
    avatars = {}
    if not args.check:
        for hid, h in heroes.items():
            png = make_avatar(hid)
            if png:
                avatars["%s/%s.png" % (IMG, hid)] = png
                h.avatar = True
    for hid, h in sorted(heroes.items()):
        wanted[h.page] = page_top(h, h.avatar) + "\n" + h.dev.rstrip("\n") + "\n"
    for k, v in imported.items():
        if not k.startswith(PAGES + "/"):
            wanted[k] = v
    readme = read_text("README.md") or ""
    table = readme_table_md(heroes)
    if BEGIN in readme and END in readme:
        a, b = readme.index(BEGIN), readme.index(END) + len(END)
        wanted["README.md"] = readme[:a] + table + readme[b:]
    else:
        print("README.md has no %s ... %s markers: the hero table is not written" % (BEGIN, END))

    stale = []
    for rel, text in sorted(wanted.items()):
        if read_text(rel) != text:
            stale.append(rel)
    pages = {h.page for h in heroes.values()}
    extra = [p for p in list_files(PAGES + "/*.md") if p not in pages]
    if args.check:
        for rel in stale:
            print("out of date: %s" % rel)
        for rel in extra:
            print("no such hero: %s" % rel)
        if stale or extra:
            print("run: python tools/docs/gen_hero_docs.py")
            return 1
        print("hero docs up to date (%d pages)" % len(pages))
        return 0
    for rel, text in wanted.items():
        if rel in stale:
            p = os.path.join(ROOT, rel)
            os.makedirs(os.path.dirname(p), exist_ok=True)
            with open(p, "w", encoding="utf-8", newline="\n") as f:
                f.write(text)
    for rel, png in avatars.items():
        p = os.path.join(ROOT, rel)
        if read_bytes(rel) != png:
            os.makedirs(os.path.dirname(p), exist_ok=True)
            with open(p, "wb") as f:
                f.write(png)
    for rel in extra:
        print("note: %s has no hero in league/champion or tft/champion" % rel)
    released = None
    if args.released:
        with open(args.released, encoding="utf-8") as f:
            released = {l.strip().replace("league_", "") for l in f if l.strip()}
    if args.xlsx:
        write_xlsx(heroes, released, args.xlsx)
    nolane = [h.id for h in heroes.values() if not h.lane]
    print("%d hero pages (%d rewritten), README table %s, %d avatars%s" % (
        len(pages), len([s for s in stale if s.startswith(PAGES)]), "updated" if "README.md" in stale else "unchanged",
        len(avatars), "; no lane: " + ", ".join(nolane) if nolane else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
